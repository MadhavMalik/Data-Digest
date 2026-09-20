"""Export a completed analysis to a human-readable folder.

The engine's internal state is JSON and Python objects.  This module turns one
run into something a person can actually open and read:

    01_profiling/          what the profiler found, and the EXACT text the LLM saw
    02_hypotheses/         what the model proposed, round by round
    03_combinations/       every variable combination tried, with its result,
                           plus what was rejected and why
    04_graphs/             every rendered graph, with an index
    05_vlm_reasoning/      the full interpretation exchange per graph:
                           context sent -> raw response -> critic verdict
    06_evidence/           the stored evidence objects
    07_summary/            final answer + telemetry

Spreadsheets are written with openpyxl when available and fall back to CSV
otherwise, so a missing optional dependency degrades the format rather than
losing the data.
"""

from __future__ import annotations

import csv
import json
import shutil
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Sequence

try:  # openpyxl is optional
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    HAVE_XLSX = True
except ImportError:  # pragma: no cover
    HAVE_XLSX = False


HEADER_FILL = "FF2A78D6"
MAX_CELL_CHARS = 32_000  # Excel's hard limit is 32,767


@dataclass
class ExportResult:
    root: Path
    files: list[Path]

    def summary(self) -> str:
        return f"{len(self.files)} files written to {self.root}"


# ---------------------------------------------------------------------------
# Spreadsheet helpers
# ---------------------------------------------------------------------------


def _clean(value: Any) -> Any:
    """Coerce a value into something a spreadsheet cell accepts."""
    if value is None or isinstance(value, (str, int, float, bool)):
        if isinstance(value, str) and len(value) > MAX_CELL_CHARS:
            return value[: MAX_CELL_CHARS - 20] + "... [truncated]"
        if isinstance(value, float) and value != value:  # NaN
            return None
        return value
    if isinstance(value, (list, tuple, set)):
        return "; ".join(str(_clean(v)) for v in value)[:MAX_CELL_CHARS]
    if isinstance(value, dict):
        return json.dumps(value, default=str)[:MAX_CELL_CHARS]
    return str(value)[:MAX_CELL_CHARS]


def write_sheet(
    path: Path,
    headers: Sequence[str],
    rows: Sequence[Sequence[Any]],
    *,
    sheet_name: str = "Sheet1",
    widths: dict[int, int] | None = None,
    wrap_columns: Sequence[int] = (),
) -> Path:
    """Write one table as .xlsx, or .csv when openpyxl is unavailable."""
    path.parent.mkdir(parents=True, exist_ok=True)

    if not HAVE_XLSX:
        csv_path = path.with_suffix(".csv")
        with csv_path.open("w", newline="") as fh:
            writer = csv.writer(fh)
            writer.writerow(headers)
            writer.writerows([[_clean(c) for c in row] for row in rows])
        return csv_path

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name[:31]

    ws.append(list(headers))
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFFFF")
        cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
        cell.alignment = Alignment(vertical="center", wrap_text=True)

    for row in rows:
        ws.append([_clean(c) for c in row])

    # Sensible default widths; explicit overrides win.
    for idx, header in enumerate(headers, start=1):
        letter = get_column_letter(idx)
        if widths and idx in widths:
            ws.column_dimensions[letter].width = widths[idx]
        else:
            longest = max(
                [len(str(header))] + [len(str(_clean(r[idx - 1]))[:60]) for r in rows[:200]]
                or [10]
            )
            ws.column_dimensions[letter].width = min(max(12, longest + 2), 60)

    for idx in wrap_columns:
        letter = get_column_letter(idx)
        ws.column_dimensions[letter].width = 70
        for cell in ws[letter]:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    ws.freeze_panes = "A2"
    wb.save(path)
    return path


def write_text(path: Path, content: str) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(content)
    return path


# ---------------------------------------------------------------------------
# The exporter
# ---------------------------------------------------------------------------


def export_analysis(result, out_root: Path, *, copy_graphs: bool = True) -> ExportResult:
    """Write a complete, readable export of one analysis run."""
    root = Path(out_root)
    root.mkdir(parents=True, exist_ok=True)
    written: list[Path] = []

    written += _export_profiling(result, root / "01_profiling")
    written += _export_hypotheses(result, root / "02_hypotheses")
    written += _export_combinations(result, root / "03_combinations")
    written += _export_graphs(result, root / "04_graphs", copy_graphs=copy_graphs)
    written += _export_vlm_reasoning(result, root / "05_vlm_reasoning")
    written += _export_evidence(result, root / "06_evidence")
    written += _export_summary(result, root / "07_summary")
    written.append(_write_index(result, root, written))

    return ExportResult(root=root, files=written)


# ---- 01 profiling ---------------------------------------------------------


def _export_profiling(result, out: Path) -> list[Path]:
    profile = result.profile
    files = []

    files.append(
        write_text(
            out / "dataset_card.txt",
            profile.dataset.to_compact_text()
            + "\n\nProfiler stats:\n"
            + json.dumps(profile.stats, indent=2),
        )
    )

    # This is the EXACT text handed to the language model. Worth reading as-is:
    # it is the whole of what the model knows about 3.7M rows.
    prompt_text = profile.to_prompt_text()
    files.append(
        write_text(
            out / "EXACT_TEXT_SENT_TO_LLM.txt",
            "This file is the complete dataset description the language model\n"
            "receives. It never sees a single row of data.\n"
            f"Size: {len(prompt_text):,} characters "
            f"(~{len(prompt_text)//4:,} tokens) standing in for "
            f"{profile.dataset.row_count:,} rows / "
            f"{profile.dataset.byte_size/1048576:.1f} MiB on disk.\n"
            + "=" * 78
            + "\n\n"
            + prompt_text,
        )
    )

    headers = [
        "column", "semantic_type", "unit", "physical_dtype", "type_confidence",
        "type_rule", "rows", "nulls", "null_%", "distinct", "min", "p25",
        "median", "p75", "max", "mean", "std", "negatives", "zeros",
        "description", "caveats", "warnings", "sample_values",
    ]
    rows = []
    for name, c in profile.columns.items():
        n = c.numeric
        rows.append([
            name, c.semantic_type.value, c.unit.label, c.physical_dtype,
            round(c.type_confidence, 3), c.type_rule, c.row_count, c.null_count,
            round(100 * c.null_fraction, 3), c.categories.distinct_count,
            n.min, n.p25, n.median, n.p75, n.max, n.mean, n.std,
            n.negative_count, n.zero_count,
            c.description, c.caveats, c.warnings, c.sample_values[:5],
        ])
    files.append(
        write_sheet(out / "column_profiles.xlsx", headers, rows,
                    sheet_name="column profiles", wrap_columns=(20, 21, 22))
    )
    return files


# ---- 02 hypotheses --------------------------------------------------------


def _export_hypotheses(result, out: Path) -> list[Path]:
    headers = [
        "round", "source", "hypothesis_id", "base_features", "target_features",
        "transformations_proposed", "relationship_types", "expected_direction",
        "priority", "est_information_gain", "est_compute_cost", "rationale",
    ]
    rows = []
    lines = ["# Hypotheses proposed by the planner", ""]

    for entry in result.state.hypotheses:
        lines.append(f"## Round {entry['round'] + 1} — source: `{entry['source']}`")
        lines.append(f"\n_{entry.get('summary', '')}_\n")
        for h in entry.get("hypotheses", []):
            rows.append([
                entry["round"] + 1, entry["source"], h.get("id", ""),
                h.get("base_features", []), h.get("target_features", []),
                h.get("transformation_candidates", []),
                h.get("relationship_types", []), h.get("expected_direction", ""),
                h.get("priority", ""), h.get("estimated_information_gain", ""),
                h.get("estimated_compute_cost", ""), h.get("rationale", ""),
            ])
            base = ", ".join(h.get("base_features", []))
            target = ", ".join(h.get("target_features", []))
            lines.append(f"- **{base} → {target}** ({h.get('priority', '')})")
            if h.get("transformation_candidates"):
                lines.append(f"  - transformations: `{'`, `'.join(h['transformation_candidates'])}`")
            if h.get("rationale"):
                lines.append(f"  - rationale: {h['rationale']}")
        lines.append("")

    files = [
        write_sheet(out / "hypotheses.xlsx", headers, rows,
                    sheet_name="hypotheses", wrap_columns=(12,)),
        write_text(out / "hypotheses.md", "\n".join(lines)),
    ]
    if result.state.dropped_hypothesis_columns:
        files.append(
            write_text(
                out / "rejected_column_references.txt",
                "Column names the model referenced that do not exist in the dataset.\n"
                "These were stripped before anything reached the data layer.\n\n"
                + "\n".join(sorted(set(result.state.dropped_hypothesis_columns))),
            )
        )
    return files


# ---- 03 combinations ------------------------------------------------------


def _export_combinations(result, out: Path) -> list[Path]:
    """Every variable combination actually tested, with its measured result."""
    headers = [
        "rank", "variable_x", "variable_y", "expression", "kind", "method",
        "n_rows", "pearson_r", "spearman_rho", "eta", "mutual_information",
        "effect_size", "direction", "strength", "stability", "r_squared",
        "ols_slope", "q_value_bh_fdr", "score", "depth", "branch",
        "equivalent_forms_collapsed", "warnings", "skipped_reason",
    ]
    rows = []
    ranked = result.state.ranked_tests(meaningful_only=False)
    for i, t in enumerate(ranked, start=1):
        r = t.result
        rows.append([
            i, r.x_name, r.y_name, t.expression_hash,
            "MECHANICAL (definitional)" if t.is_mechanical else "empirical",
            r.method, r.n, r.pearson_r, r.spearman_rho, r.eta,
            r.mutual_information, round(r.effect, 5), r.direction,
            r.strength_label(), r.stability, r.r_squared, r.slope, r.q_value,
            round(t.score.total, 5), t.depth, t.branch_id,
            r.extra.get("equivalent_forms", []), r.warnings, r.skipped_reason,
        ])

    files = [
        write_sheet(out / "ALL_COMBINATIONS_TESTED.xlsx", headers, rows,
                    sheet_name="combinations tested", wrap_columns=(22, 23))
    ]

    # What the dimensional layer threw away, and why.
    prune = result.prune_stats
    pr_headers = ["rejection_reason", "candidates_rejected"]
    pr_rows = sorted(prune.reasons.items(), key=lambda kv: -kv[1])
    files.append(
        write_sheet(out / "candidates_rejected_by_reason.xlsx", pr_headers, pr_rows,
                    sheet_name="pruning", widths={1: 55, 2: 22})
    )

    files.append(
        write_text(
            out / "pruning_summary.txt",
            "CANDIDATE PRUNING\n"
            + "=" * 78
            + f"""
Candidates considered          : {prune.considered:,}
Emitted for testing            : {prune.emitted:,}
Rejected by dimensional units  : {prune.pruned_by_units:,}
Rejected by policy             : {prune.pruned_by_policy:,}
Rejected as duplicates         : {prune.pruned_by_dedup:,}
Rejected by budget             : {prune.pruned_by_budget:,}
Overall prune rate             : {prune.prune_rate:.1%}

Rejection happens on UNIT ALGEBRA before any row of data is read.
`dollars + miles` and `zone_id * fare` are not merely low-scoring --
they are never constructed, so they cannot be reported.

Breakdown by reason:
"""
            + "\n".join(f"  {reason:<52} {count:>8,}" for reason, count in pr_rows),
        )
    )
    return files


# ---- 04 graphs ------------------------------------------------------------


def _export_graphs(result, out: Path, *, copy_graphs: bool) -> list[Path]:
    out.mkdir(parents=True, exist_ok=True)
    files = []
    headers = [
        "graph_file", "plot_type", "variable_x", "variable_y", "expression",
        "n_rows_represented", "rendering_strategy", "statistics_shown",
        "evidence_id", "explanation_status",
    ]
    rows = []

    for e in result.evidence:
        if not e.plot_uri:
            continue
        src = Path(e.plot_uri)
        name = src.name
        if copy_graphs and src.exists():
            shutil.copy2(src, out / name)
        rows.append([
            name, e.plot_type,
            e.feature_names[0] if e.feature_names else "",
            e.feature_names[1] if len(e.feature_names) > 1 else "",
            e.canonical_expression, e.sample_size, e.plot_description,
            e.statistical_metrics.to_compact_text(), e.evidence_id,
            e.explanation_status.value,
        ])

    files.append(
        write_sheet(out / "graph_index.xlsx", headers, rows,
                    sheet_name="graphs", wrap_columns=(7, 8))
    )
    return files


# ---- 05 VLM reasoning (the headline deliverable) --------------------------


def _export_vlm_reasoning(result, out: Path) -> list[Path]:
    """The full second-stage exchange: what the model was shown, what it said."""
    traces = result.vlm_traces
    files = []

    # --- the readable narrative version ---
    lines = [
        "# Stage 2 — Multimodal interpretation of the graphs",
        "",
        f"Analysis `{result.analysis_id}` · {len(traces)} graph interpretations",
        "",
        "For each graph the engine rendered, this records **exactly what the model",
        "received** (the image, the statistics, the column definitions, the filters,",
        "the caveats) and **exactly what it returned**, followed by the deterministic",
        "critic's verdict on whether the claim is consistent with the numbers.",
        "",
        "---",
        "",
    ]

    for t in traces:
        interp = t.get("interpretation_after_critic") or t.get("interpretation") or {}
        critic = t.get("critic", {})

        lines += [
            f"## {t['sequence']}. `{t['expression']}` vs `{t['y']}`",
            "",
            f"**Graph:** `{t.get('plot_type')}` → `{Path(t['plot_path']).name if t.get('plot_path') else 'not rendered'}`",
            f"**Image actually sent to the model:** {'YES' if t.get('image_sent') else 'NO'}"
            + (f" ({t.get('image_bytes', 0):,} bytes)" if t.get("image_sent") else ""),
            f"**Model:** `{t.get('model') or t.get('source')}`"
            + (f" · {t.get('prompt_tokens', 0):,} prompt + {t.get('completion_tokens', 0):,} completion tokens"
               f" · {t.get('latency_seconds', 0):.2f}s" if t.get("model") else ""),
            "",
            "### What the model was told (verbatim)",
            "",
            "**Statistics (authoritative):**",
            "```",
            t.get("prompt_statistics", ""),
            "```",
            "",
            "**Column definitions:**",
            "```",
            t.get("prompt_column_context", ""),
            "```",
            "",
            "**How the graph was rendered:**",
            "```",
            t.get("plot_description_sent", ""),
            "```",
            "",
        ]
        if t.get("prompt_caveats"):
            lines += ["**Caveats supplied:**", ""]
            lines += [f"- {c}" for c in t["prompt_caveats"]]
            lines.append("")
        if t.get("prompt_related_evidence"):
            lines += ["**Related prior evidence retrieved from memory:**", ""]
            lines += [f"- {e}" for e in t["prompt_related_evidence"]]
            lines.append("")

        lines += ["### What the model concluded", ""]
        if interp.get("observation"):
            lines += ["**Observation:**", "", interp["observation"], ""]
        if interp.get("interpretation"):
            lines += ["**Interpretation:**", "", interp["interpretation"], ""]
        for key, label in (
            ("plausible_mechanisms", "Plausible mechanisms"),
            ("alternative_explanations", "Alternative explanations"),
            ("confounders", "Confounders"),
            ("additional_tests", "Tests it suggested"),
        ):
            if interp.get(key):
                lines += [f"**{label}:**", ""]
                lines += [f"- {v}" for v in interp[key]]
                lines.append("")
        lines += [
            f"**Stated direction:** `{interp.get('stated_direction')}` · "
            f"**Strength:** `{interp.get('strength')}` · "
            f"**Status:** `{interp.get('conclusion_status')}` · "
            f"**Confidence:** {interp.get('confidence')}",
            "",
            "### Critic verdict",
            "",
        ]
        if critic.get("violations"):
            for v in critic["violations"]:
                lines.append(f"- `{v['severity'].upper()}` **{v['check']}** — {v['message']}")
                if v.get("suggested_fix"):
                    lines.append(f"  - suggested fix: {v['suggested_fix']}")
        else:
            lines.append("- No issues. The claim is consistent with the statistics, the units,")
            lines.append("  and the dataset's documented semantics.")
        lines += ["", "---", ""]

    files.append(write_text(out / "VLM_REASONING.md", "\n".join(lines)))

    # --- the spreadsheet version ---
    headers = [
        "seq", "expression", "vs", "plot_type", "image_sent", "model",
        "observation", "interpretation", "plausible_mechanisms",
        "alternative_explanations", "confounders", "stated_direction",
        "measured_direction", "strength_claimed", "conclusion_status",
        "confidence", "critic_passed", "critic_errors", "critic_warnings",
        "critic_detail", "prompt_tokens", "completion_tokens", "latency_s",
    ]
    rows = []
    for t in traces:
        interp = t.get("interpretation_after_critic") or t.get("interpretation") or {}
        critic = t.get("critic", {})
        measured = ""
        for line in (t.get("prompt_statistics") or "").splitlines():
            if line.startswith("DIRECTION AS MEASURED:"):
                measured = line.split(":", 1)[1].strip()
        rows.append([
            t["sequence"], t["expression"], t["y"], t.get("plot_type"),
            t.get("image_sent"), t.get("model") or t.get("source"),
            interp.get("observation"), interp.get("interpretation"),
            interp.get("plausible_mechanisms"), interp.get("alternative_explanations"),
            interp.get("confounders"), interp.get("stated_direction"), measured,
            interp.get("strength"), interp.get("conclusion_status"),
            interp.get("confidence"), critic.get("passed"),
            critic.get("error_count"), critic.get("warning_count"),
            [f"{v['severity']}:{v['check']}" for v in critic.get("violations", [])],
            t.get("prompt_tokens"), t.get("completion_tokens"), t.get("latency_seconds"),
        ])
    files.append(
        write_sheet(out / "vlm_reasoning.xlsx", headers, rows,
                    sheet_name="VLM reasoning", wrap_columns=(7, 8, 9, 10, 11, 20))
    )

    # --- raw, unparsed model output, for auditing ---
    raw = {
        t["sequence"]: {
            "expression": t["expression"],
            "raw_response": t.get("raw_response"),
            "source": t.get("source"),
            "fallback_reason": t.get("fallback_reason"),
        }
        for t in traces
    }
    files.append(
        write_text(out / "raw_model_responses.json", json.dumps(raw, indent=2, default=str))
    )
    return files


# ---- 06 evidence ----------------------------------------------------------


def _export_evidence(result, out: Path) -> list[Path]:
    headers = [
        "evidence_id", "expression", "features", "status", "confidence",
        "n", "pearson_r", "spearman_rho", "eta", "mutual_information",
        "effect", "direction", "strength", "stability", "plot_type",
        "plot_file", "summary", "tags", "related_evidence", "resolved_by",
        "warnings", "filters_applied",
    ]
    rows = []
    for e in result.evidence:
        m = e.statistical_metrics
        rows.append([
            e.evidence_id, e.canonical_expression, e.feature_names,
            e.explanation_status.value, round(e.explanation_confidence, 3),
            m.n, m.pearson_r, m.spearman_rho, m.eta, m.mutual_information,
            round(m.effect, 5), m.direction, m.strength, m.stability,
            e.plot_type, Path(e.plot_uri).name if e.plot_uri else "",
            e.textual_summary, e.tags, e.related_evidence_ids, e.resolved_by_ids,
            e.warnings, e.filters,
        ])
    files = [
        write_sheet(out / "evidence.xlsx", headers, rows,
                    sheet_name="evidence", wrap_columns=(17, 21))
    ]
    files.append(
        write_text(
            out / "evidence_full.json",
            json.dumps([e.to_dict() for e in result.evidence], indent=2, default=str),
        )
    )
    return files


# ---- 07 summary -----------------------------------------------------------


def _export_summary(result, out: Path) -> list[Path]:
    metrics = result.metrics.to_dict() if result.metrics else {}
    counters = metrics.get("counters", {})
    derived = metrics.get("derived", {})
    view = result.view_report
    ds = result.profile.dataset

    lines = [
        f"# {result.analysis_id} — final answer",
        "",
        f"**Question:** {result.question}",
        f"**Dataset:** {ds.dataset_id} — {ds.row_count:,} rows × {ds.column_count} columns",
        f"**Stop reason:** `{result.state.stop_reason.value if result.state.stop_reason else 'n/a'}`",
        "",
    ]
    if result.final_answer:
        fa = result.final_answer
        lines += ["## Answer", "", fa.answer, ""]
        if fa.key_findings:
            lines += ["### Key findings", ""] + [f"- {f}" for f in fa.key_findings] + [""]
        if fa.caveats:
            lines += ["### Caveats", ""] + [f"- {c}" for c in fa.caveats] + [""]
        if fa.unresolved_questions:
            lines += ["### Unresolved", ""] + [f"- {u}" for u in fa.unresolved_questions] + [""]
    if result.degradations:
        lines += ["## Degradations", ""] + [f"- {d}" for d in result.degradations] + [""]

    files = [write_text(out / "final_answer.md", "\n".join(lines))]

    metric_rows = [
        ("Raw dataset rows", ds.row_count),
        ("Raw bytes on disk", ds.byte_size),
        ("Columns profiled", counters.get("columns_profiled", 0)),
        ("Rows after validity filters", view.get("rows_after", 0)),
        ("Rows excluded", view.get("rows_excluded", 0)),
        ("Exclusion fraction", round(view.get("exclusion_fraction", 0), 5)),
        ("Candidate expressions considered", result.prune_stats.considered),
        ("Pruned by dimensional analysis", result.prune_stats.pruned_by_units),
        ("Pruned by policy", result.prune_stats.pruned_by_policy),
        ("Prune rate", round(result.prune_stats.prune_rate, 4)),
        ("Statistical tests run", counters.get("statistical_tests", 0)),
        ("Mutual-information tests", counters.get("mutual_information_tests", 0)),
        ("Stability tests", counters.get("stability_tests", 0)),
        ("Group comparisons", counters.get("group_comparisons", 0)),
        ("Graphs rendered", counters.get("plots_rendered", 0)),
        ("LLM calls (planning)", counters.get("llm_calls", 0)),
        ("VLM calls (graph reading)", counters.get("vlm_calls", 0)),
        ("Model cache hits", derived.get("model_cache_hits", 0)),
        ("Model cache hit rate", derived.get("model_cache_hit_rate", 0)),
        ("Total prompt tokens", derived.get("total_prompt_tokens", 0)),
        ("Total completion tokens", derived.get("total_completion_tokens", 0)),
        ("Tokens avoided by cache", derived.get("tokens_avoided_by_cache", 0)),
        ("Evidence objects created", counters.get("evidence_created", 0)),
        ("Prior evidence retrieved", counters.get("evidence_retrievals", 0)),
        ("Joint reinterpretations", counters.get("joint_reinterpretations", 0)),
        ("Critic errors caught", counters.get("critic_errors", 0)),
        ("Critic warnings", counters.get("critic_warnings", 0)),
        ("Rejected model transformations", counters.get("rejected_transformations", 0)),
        ("Wall-clock seconds", metrics.get("elapsed_seconds", 0)),
    ]
    files.append(
        write_sheet(out / "metrics.xlsx", ["metric", "value"], metric_rows,
                    sheet_name="metrics", widths={1: 42, 2: 20})
    )

    filter_rows = [
        (name, view.get("excluded_by_rule", {}).get(name, 0), rationale)
        for name, rationale in view.get("rationale", {}).items()
    ]
    files.append(
        write_sheet(out / "analysis_view_filters.xlsx",
                    ["filter", "rows_excluded", "rationale"], filter_rows,
                    sheet_name="filters", wrap_columns=(3,))
    )
    files.append(
        write_text(out / "full_result.json", json.dumps(result.to_dict(), indent=2, default=str))
    )
    return files


# ---- index ----------------------------------------------------------------


def _write_index(result, root: Path, written: list[Path]) -> Path:
    ds = result.profile.dataset
    metrics = result.metrics.to_dict() if result.metrics else {}
    counters = metrics.get("counters", {})

    lines = [
        "# Analysis output",
        "",
        f"**Question:** {result.question}",
        f"**Dataset:** {ds.dataset_id} — {ds.row_count:,} rows × {ds.column_count} columns "
        f"({ds.byte_size/1048576:.1f} MiB)",
        f"**Analysis ID:** `{result.analysis_id}`",
        f"**Generated:** {datetime.now(timezone.utc).isoformat(timespec='seconds')}",
        "",
        "## Where to look",
        "",
        "| Folder | What's in it |",
        "|---|---|",
        "| `01_profiling/` | What the profiler found. **`EXACT_TEXT_SENT_TO_LLM.txt`** is the complete dataset description the model receives — it never sees a row. |",
        "| `02_hypotheses/` | What the planner proposed each round, and any hallucinated column names that were stripped. |",
        "| `03_combinations/` | **`ALL_COMBINATIONS_TESTED.xlsx`** — every variable combination tried with its measured result. Plus what was rejected and why. |",
        "| `04_graphs/` | Every rendered graph, with an index. |",
        "| `05_vlm_reasoning/` | **`VLM_REASONING.md`** — stage 2. What the model was shown and what it concluded, per graph, with the critic's verdict. |",
        "| `06_evidence/` | The stored evidence objects. |",
        "| `07_summary/` | Final answer, telemetry, filter accounting. |",
        "",
        "## Headline numbers",
        "",
        f"- Candidate expressions considered: **{result.prune_stats.considered:,}**",
        f"- Pruned by dimensional analysis: **{result.prune_stats.pruned_by_units:,}** ({result.prune_stats.prune_rate:.0%} overall prune rate)",
        f"- Statistical tests run: **{counters.get('statistical_tests', 0):,}**",
        f"- Graphs rendered: **{counters.get('plots_rendered', 0)}**",
        f"- Planning (LLM) calls: **{counters.get('llm_calls', 0)}**",
        f"- Graph-reading (VLM) calls: **{counters.get('vlm_calls', 0)}**",
        f"- Critic errors caught: **{counters.get('critic_errors', 0)}**",
        f"- Wall-clock: **{metrics.get('elapsed_seconds', 0):.1f}s**",
        "",
        f"{len(written)} files written.",
    ]
    return write_text(root / "README.md", "\n".join(lines))
